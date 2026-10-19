import sys
from docx import Document
from docx.shared import Pt
p = "/tmp/template-ref.docx"
d = Document(p)
print("SECTION page:", d.sections[0].page_width.cm, d.sections[0].page_height.cm,
      "margins:", d.sections[0].left_margin.cm, d.sections[0].right_margin.cm,
      d.sections[0].top_margin.cm, d.sections[0].bottom_margin.cm)
print("paras:", len(d.paragraphs), "tables:", len(d.tables))
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if not t: continue
    runs = []
    for r in para.runs[:4]:
        rpr = r._element.rPr
        ea = None
        if rpr is not None and rpr.rFonts is not None:
            ea = rpr.rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia')
        runs.append((r.text[:20], r.font.name, ea, r.font.size.pt if r.font.size else None, r.bold))
    print(i, "|", para.style.name, "|", para.alignment, "|", t[:60], "|", runs)
